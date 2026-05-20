from rest_framework import viewsets, status, permissions
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated, AllowAny
from django.db.models import Count, Avg
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
import io
from datetime import datetime

from .models import Survey, Question, Option, SurveyResponse, Answer
from .serializers import (
    SurveyListSerializer, SurveySerializer, SurveyPublicSerializer,
    SurveyResponseCreateSerializer, SurveyResponseDetailSerializer,
)


class IsOwnerOrReadOnly(permissions.BasePermission):
    def has_object_permission(self, request, view, obj):
        if request.method in permissions.SAFE_METHODS:
            return True
        return obj.creator == request.user or request.user.is_staff


class SurveyViewSet(viewsets.ModelViewSet):
    queryset = Survey.objects.all()
    permission_classes = [IsAuthenticated, IsOwnerOrReadOnly]

    def get_serializer_class(self):
        if self.action == 'list':
            return SurveyListSerializer
        return SurveySerializer

    def get_queryset(self):
        user = self.request.user
        if not user.is_authenticated:
            return Survey.objects.filter(status='published')
        if user.is_staff:
            return Survey.objects.all()
        return Survey.objects.filter(creator=user)

    def perform_create(self, serializer):
        serializer.save(creator=self.request.user)

    @action(detail=True, methods=['post'], permission_classes=[IsAuthenticated])
    def publish(self, request, pk=None):
        survey = self.get_object()
        if survey.creator != request.user and not request.user.is_staff:
            return Response({'detail': '没有权限'}, status=status.HTTP_403_FORBIDDEN)
        survey.publish()
        from django.conf import settings
        survey_url = f"{settings.FRONTEND_URL}/survey/{survey.id}"
        return Response({
            'status': 'published',
            'qr_code': survey.qr_code.url if survey.qr_code else None,
            'survey_url': survey_url
        })

    @action(detail=True, methods=['post'], permission_classes=[IsAuthenticated])
    def close(self, request, pk=None):
        survey = self.get_object()
        if survey.creator != request.user and not request.user.is_staff:
            return Response({'detail': '没有权限'}, status=status.HTTP_403_FORBIDDEN)
        survey.status = 'closed'
        survey.save()
        return Response({'status': 'closed'})

    @action(detail=True, methods=['get'], permission_classes=[AllowAny])
    def public(self, request, pk=None):
        survey = self.get_object()
        if survey.status != 'published':
            return Response({'detail': '问卷未发布'}, status=status.HTTP_404_NOT_FOUND)
        serializer = SurveyPublicSerializer(survey)
        return Response(serializer.data)

    @action(detail=True, methods=['post'], permission_classes=[AllowAny])
    def submit(self, request, pk=None):
        survey = self.get_object()
        if survey.status != 'published':
            return Response({'detail': '问卷未发布'}, status=status.HTTP_400_BAD_REQUEST)
        
        ip_address = request.META.get('HTTP_X_FORWARDED_FOR', request.META.get('REMOTE_ADDR', ''))
        respondent = request.user if request.user.is_authenticated else None

        serializer = SurveyResponseCreateSerializer(
            data=request.data,
            context={'survey_id': pk, 'ip_address': ip_address, 'respondent': respondent}
        )
        serializer.is_valid(raise_exception=True)
        serializer.save()
        return Response({'status': 'success'}, status=status.HTTP_201_CREATED)

    @action(detail=True, methods=['get'], permission_classes=[IsAuthenticated])
    def responses(self, request, pk=None):
        survey = self.get_object()
        if survey.creator != request.user and not request.user.is_staff:
            return Response({'detail': '没有权限'}, status=status.HTTP_403_FORBIDDEN)
        responses = survey.responses.filter(completed=True)
        serializer = SurveyResponseDetailSerializer(responses, many=True)
        return Response(serializer.data)

    @action(detail=True, methods=['get'], permission_classes=[IsAuthenticated])
    def stats(self, request, pk=None):
        survey = self.get_object()
        if survey.creator != request.user and not request.user.is_staff:
            return Response({'detail': '没有权限'}, status=status.HTTP_403_FORBIDDEN)
        
        stats = []
        for question in survey.questions.all():
            question_stat = {
                'question_id': question.id,
                'question_title': question.title,
                'question_type': question.question_type,
                'stats': {}
            }
            
            answers = Answer.objects.filter(response__survey=survey, question=question, response__completed=True)
            total_answers = answers.count()
            question_stat['stats']['total'] = total_answers
            
            if question.question_type == 'single_choice':
                option_counts = answers.values('selected_options__text').annotate(count=Count('id')).order_by('-count')
                question_stat['stats']['options'] = [
                    {'label': item['selected_options__text'] or '未填写', 'value': item['count']}
                    for item in option_counts if item['selected_options__text']
                ]
                
            elif question.question_type == 'multiple_choice':
                option_counts = {}
                for answer in answers:
                    for option in answer.selected_options.all():
                        option_counts[option.text] = option_counts.get(option.text, 0) + 1
                question_stat['stats']['options'] = [
                    {'label': label, 'value': count}
                    for label, count in sorted(option_counts.items(), key=lambda x: -x[1])
                ]
                
            elif question.question_type == 'rating':
                avg_rating = answers.aggregate(avg=Avg('rating_value'))['avg'] or 0
                rating_dist = answers.values('rating_value').annotate(count=Count('id')).order_by('rating_value')
                question_stat['stats']['average'] = round(avg_rating, 2)
                question_stat['stats']['distribution'] = [
                    {'rating': item['rating_value'], 'count': item['count']}
                    for item in rating_dist if item['rating_value']
                ]
                
            elif question.question_type == 'text':
                question_stat['stats']['answers'] = [
                    answer.text_answer for answer in answers if answer.text_answer
                ]
            
            stats.append(question_stat)
        
        return Response(stats)

    @action(detail=True, methods=['get'], permission_classes=[IsAuthenticated])
    def export(self, request, pk=None):
        survey = self.get_object()
        if survey.creator != request.user and not request.user.is_staff:
            return Response({'detail': '没有权限'}, status=status.HTTP_403_FORBIDDEN)
        
        wb = Workbook()
        ws = wb.active
        ws.title = '问卷结果'
        
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')
        
        headers = ['回答ID', '回答者', '邮箱', '开始时间', '完成时间']
        question_map = {}
        for idx, question in enumerate(survey.questions.all(), 1):
            header = f"Q{idx}. {question.title[:20]}"
            headers.append(header)
            question_map[question.id] = idx
        
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
        
        responses = survey.responses.filter(completed=True)
        for row_idx, response in enumerate(responses, 2):
            ws.cell(row=row_idx, column=1, value=str(response.id)[:8])
            ws.cell(row=row_idx, column=2, value=response.respondent_name or '匿名')
            ws.cell(row=row_idx, column=3, value=response.respondent_email or '-')
            ws.cell(row=row_idx, column=4, value=response.started_at.strftime('%Y-%m-%d %H:%M:%S') if response.started_at else '-')
            ws.cell(row=row_idx, column=5, value=response.completed_at.strftime('%Y-%m-%d %H:%M:%S') if response.completed_at else '-')
            
            for answer in response.answers.all():
                col_idx = 4 + question_map.get(answer.question.id, 0)
                if answer.question.question_type in ['single_choice', 'multiple_choice']:
                    value = ', '.join([opt.text for opt in answer.selected_options.all()])
                elif answer.question.question_type == 'rating':
                    value = answer.rating_value or ''
                else:
                    value = answer.text_answer or ''
                ws.cell(row=row_idx, column=col_idx, value=value)
        
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[column].width = min(max_length + 2, 50)
        
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(output, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        from urllib.parse import quote
        filename = f'{survey.title}_{datetime.now().strftime("%Y%m%d")}.xlsx'
        response['Content-Disposition'] = f"attachment; filename*=UTF-8''{quote(filename)}"
        return response


class AdminStatsView(viewsets.ViewSet):
    permission_classes = [IsAuthenticated]

    @action(detail=False, methods=['get'])
    def overview(self, request):
        if not request.user.is_staff:
            return Response({'detail': '没有权限'}, status=status.HTTP_403_FORBIDDEN)
        
        total_surveys = Survey.objects.count()
        published_surveys = Survey.objects.filter(status='published').count()
        total_responses = SurveyResponse.objects.filter(completed=True).count()
        
        surveys = Survey.objects.all()
        survey_stats = []
        for survey in surveys:
            survey_stats.append({
                'survey_id': survey.id,
                'title': survey.title,
                'creator': survey.creator.username,
                'status': survey.status,
                'response_count': survey.response_count,
                'average_completion_time': survey.average_completion_time,
                'created_at': survey.created_at,
                'published_at': survey.published_at,
            })
        
        return Response({
            'total_surveys': total_surveys,
            'published_surveys': published_surveys,
            'total_responses': total_responses,
            'surveys': survey_stats,
        })
