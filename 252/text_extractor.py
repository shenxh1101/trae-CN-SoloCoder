import pdfplumber
from docx import Document
import os
from PyPDF2 import PdfReader

try:
    from pdfminer.high_level import extract_text as pdfminer_extract
    from pdfminer.layout import LAParams
    HAS_PDFMINER = True
except ImportError:
    HAS_PDFMINER = False

def extract_text_from_pdf(pdf_path):
    text = ""
    format_issues = []
    
    try:
        with pdfplumber.open(pdf_path) as pdf:
            total_pages = len(pdf.pages)
            for page_num, page in enumerate(pdf.pages):
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
                
                tables = page.find_tables()
                if tables:
                    format_issues.append(f"第{page_num + 1}页包含表格，ATS系统可能无法正确解析")
                
                images = page.images
                if images:
                    format_issues.append(f"第{page_num + 1}页包含图片，ATS系统可能无法提取图片中的文字")
    except Exception as e:
        format_issues.append(f"pdfplumber解析警告: {str(e)}")
    
    if not text.strip():
        try:
            reader = PdfReader(pdf_path)
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        except Exception as e:
            format_issues.append(f"PyPDF2解析警告: {str(e)}")
    
    if not text.strip() and HAS_PDFMINER:
        try:
            text = pdfminer_extract(pdf_path, laparams=LAParams())
        except Exception as e:
            format_issues.append(f"pdfminer解析警告: {str(e)}")
    
    if not text.strip():
        format_issues.append("PDF可能是扫描件或图片格式，无法提取文本，建议使用OCR工具处理后重新上传")
            
    return text.strip(), format_issues

def extract_text_from_docx(docx_path):
    text = ""
    format_issues = []
    
    try:
        doc = Document(docx_path)
        
        for table in doc.tables:
            format_issues.append("文档包含表格，ATS系统可能无法正确解析")
            break
        
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref:
                format_issues.append("文档包含图片，ATS系统可能无法提取图片中的文字")
                break
        
        for para in doc.paragraphs:
            if para.text.strip():
                text += para.text + "\n"
        
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text += cell.text + "\n"
                        
    except Exception as e:
        format_issues.append(f"DOCX解析错误: {str(e)}")
    
    return text.strip(), format_issues

def extract_text_from_txt(txt_path):
    text = ""
    format_issues = []
    
    try:
        with open(txt_path, 'r', encoding='utf-8') as f:
            text = f.read()
    except UnicodeDecodeError:
        try:
            with open(txt_path, 'r', encoding='gbk') as f:
                text = f.read()
        except Exception as e:
            format_issues.append(f"文本文件编码错误: {str(e)}")
    except Exception as e:
        format_issues.append(f"文本文件读取错误: {str(e)}")
    
    return text.strip(), format_issues

def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    
    if ext == '.pdf':
        return extract_text_from_pdf(file_path)
    elif ext == '.docx':
        return extract_text_from_docx(file_path)
    elif ext == '.txt':
        return extract_text_from_txt(file_path)
    else:
        return "", [f"不支持的文件格式: {ext}"]

def allowed_file(filename, allowed_extensions):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in allowed_extensions
