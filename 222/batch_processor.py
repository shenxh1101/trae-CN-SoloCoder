import os
import glob
from typing import Dict, List, Any
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
from parser import MeetingParser
from llm_client import LLMClient
from summarizer import MeetingSummarizer, SimpleSummarizer, merge_summary_with_stats
from output_formatter import OutputFormatter

try:
    from openpyxl import Workbook
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class BatchProcessor:
    def __init__(self, llm_client: LLMClient = None, use_llm: bool = True):
        self.parser = MeetingParser()
        self.llm_client = llm_client
        self.use_llm = use_llm and (llm_client is not None and llm_client.is_available())
        
        if self.use_llm:
            self.summarizer = MeetingSummarizer(llm_client)
        else:
            self.simple_summarizer = SimpleSummarizer()
    
    def process_directory(self, input_dir: str, output_dir: str = None, 
                         output_format: str = "markdown",
                         workers: int = 1) -> List[Dict[str, Any]]:
        if not os.path.exists(input_dir):
            raise ValueError(f"目录不存在: {input_dir}")
        
        txt_files = glob.glob(os.path.join(input_dir, "*.txt"))
        
        if not txt_files:
            print(f"在 {input_dir} 中未找到 .txt 文件")
            return []
        
        print(f"找到 {len(txt_files)} 个会议记录文件")
        if workers > 1:
            print(f"并发模式: {workers} 个工作线程")
        
        results = []
        
        if workers > 1:
            with ThreadPoolExecutor(max_workers=workers) as executor:
                future_to_file = {
                    executor.submit(self.process_single_file, file_path): file_path 
                    for file_path in txt_files
                }
                
                for idx, future in enumerate(as_completed(future_to_file), 1):
                    file_path = future_to_file[future]
                    filename = os.path.basename(file_path)
                    print(f"\n处理文件 [{idx}/{len(txt_files)}]: {filename}")
                    
                    try:
                        result = future.result()
                        results.append(result)
                        
                        if output_dir:
                            self._save_single_result(result, output_dir, output_format)
                            print(f"  ✓ 已保存")
                    except Exception as e:
                        print(f"  ✗ 处理失败: {str(e)}")
        else:
            for idx, file_path in enumerate(txt_files, 1):
                print(f"\n处理文件 [{idx}/{len(txt_files)}]: {os.path.basename(file_path)}")
                
                try:
                    result = self.process_single_file(file_path)
                    results.append(result)
                    
                    if output_dir:
                        self._save_single_result(result, output_dir, output_format)
                        print(f"  ✓ 已保存")
                except Exception as e:
                    print(f"  ✗ 处理失败: {str(e)}")
        
        return results
    
    def process_single_file(self, file_path: str) -> Dict[str, Any]:
        utterances = self.parser.parse_file(file_path)
        
        stats = self.parser.get_speaker_stats(utterances)
        metadata = self.parser.get_meeting_metadata(utterances, os.path.basename(file_path))
        
        if self.use_llm:
            summary = self.summarizer.execute(utterances, self.parser)
        else:
            summary = {
                "summary_by_speaker": {},
                "meeting_consensus": [],
                "todo_items": self.simple_summarizer.extract_todo_simple(utterances),
                "time_events": self.simple_summarizer.extract_time_events_simple(utterances),
                "overall_summary": "(无AI总结，仅使用关键词提取模式)"
            }
        
        result = merge_summary_with_stats(summary, stats, metadata)
        return result
    
    def _save_single_result(self, result: Dict, output_dir: str, output_format: str):
        if not os.path.exists(output_dir):
            os.makedirs(output_dir)
        
        filename = result.get("metadata", {}).get("filename", "unknown")
        basename = os.path.splitext(filename)[0]
        
        if output_format == "json":
            content = OutputFormatter.to_json(result)
            ext = ".json"
            is_binary = False
        elif output_format == "excel":
            ext = ".xlsx"
            is_binary = True
        elif output_format == "markdown":
            content = OutputFormatter.to_markdown(result)
            ext = ".md"
            is_binary = False
        else:
            content = OutputFormatter.to_plain_text(result)
            ext = ".txt"
            is_binary = False
        
        output_path = os.path.join(output_dir, f"{basename}_summary{ext}")
        
        if is_binary:
            OutputFormatter.to_excel(result, output_path)
        else:
            OutputFormatter.save_to_file(content, output_path)
    
    def generate_comparison_report(self, results: List[Dict[str, Any]], 
                                   output_file: str, 
                                   format: str = "markdown") -> str:
        if format == "json":
            report = self._generate_comparison_json(results)
            content = OutputFormatter.to_json(report)
            OutputFormatter.save_to_file(content, output_file)
        elif format == "excel":
            self._generate_comparison_excel(results, output_file)
        else:
            content = self._generate_comparison_markdown(results)
            OutputFormatter.save_to_file(content, output_file)
        
        return output_file
    
    def _generate_comparison_excel(self, results: List[Dict[str, Any]], output_file: str):
        if not OPENPYXL_AVAILABLE:
            raise ImportError("请安装 openpyxl: pip install openpyxl")
        
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        
        wb = Workbook()
        
        header_font = Font(bold=True, size=11, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center")
        wrap_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        
        def style_header(ws, row=1):
            for cell in ws[row]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border
        
        ws1 = wb.active
        ws1.title = "效率对比"
        ws1.append(["会议文件", "参会人数", "总发言次数", "总字数", "预估时长(分钟)", "平均发言长度"])
        
        for result in results:
            metadata = result.get("metadata", {})
            avg_length = 0
            if metadata.get("total_utterances", 0) > 0:
                avg_length = round(metadata.get("total_words", 0) / metadata.get("total_utterances", 1), 1)
            
            ws1.append([
                metadata.get('filename', 'N/A'),
                metadata.get('speaker_count', 0),
                metadata.get('total_utterances', 0),
                metadata.get('total_words', 0),
                metadata.get('estimated_duration_minutes', 0),
                avg_length
            ])
        
        style_header(ws1)
        for col in ['A', 'B', 'C', 'D', 'E', 'F']:
            ws1.column_dimensions[col].width = 18
        for row in ws1.iter_rows(min_row=1, max_row=ws1.max_row, min_col=1, max_col=6):
            for cell in row:
                cell.border = thin_border
                cell.alignment = center_align
        
        all_speakers = set()
        for result in results:
            stats = result.get("speaker_statistics", {})
            all_speakers.update(stats.keys())
        
        ws2 = wb.create_sheet("活跃度对比")
        headers = ["发言人"] + [r.get("metadata", {}).get("filename", "N/A") for r in results]
        ws2.append(headers)
        
        for speaker in sorted(all_speakers):
            row = [speaker]
            for result in results:
                stats = result.get("speaker_statistics", {})
                if speaker in stats:
                    row.append(stats[speaker].get("utterance_count", 0))
                else:
                    row.append(0)
            ws2.append(row)
        
        style_header(ws2)
        ws2.column_dimensions['A'].width = 15
        for col_idx in range(2, len(headers) + 1):
            ws2.column_dimensions[chr(64 + col_idx)].width = 20
        for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row, min_col=1, max_col=len(headers)):
            for cell in row:
                cell.border = thin_border
                cell.alignment = center_align
        
        ws3 = wb.create_sheet("待办汇总")
        ws3.append(["来源会议", "任务", "负责人", "截止时间"])
        
        all_todos = []
        for result in results:
            filename = result.get("metadata", {}).get("filename", "N/A")
            todos = result.get("todo_items", [])
            for todo in todos:
                ws3.append([
                    filename,
                    todo.get('task', ''),
                    todo.get('assignee', ''),
                    todo.get('deadline', '')
                ])
        
        style_header(ws3)
        ws3.column_dimensions['A'].width = 25
        ws3.column_dimensions['B'].width = 50
        ws3.column_dimensions['C'].width = 12
        ws3.column_dimensions['D'].width = 18
        for row in ws3.iter_rows(min_row=1, max_row=ws3.max_row, min_col=1, max_col=4):
            for cell in row:
                cell.border = thin_border
                if cell.row > 1:
                    cell.alignment = wrap_align
        
        ws4 = wb.create_sheet("时间节点")
        ws4.append(["来源会议", "事件", "时间"])
        
        for result in results:
            filename = result.get("metadata", {}).get("filename", "N/A")
            events = result.get("time_events", [])
            for event in events:
                ws4.append([
                    filename,
                    event.get('event', ''),
                    event.get('time', '')
                ])
        
        style_header(ws4)
        ws4.column_dimensions['A'].width = 25
        ws4.column_dimensions['B'].width = 50
        ws4.column_dimensions['C'].width = 18
        for row in ws4.iter_rows(min_row=1, max_row=ws4.max_row, min_col=1, max_col=3):
            for cell in row:
                cell.border = thin_border
                if cell.row > 1:
                    cell.alignment = wrap_align
        
        dir_path = os.path.dirname(output_file)
        if dir_path and not os.path.exists(dir_path):
            os.makedirs(dir_path, exist_ok=True)
        
        wb.save(output_file)
    
    def _generate_comparison_markdown(self, results: List[Dict[str, Any]]) -> str:
        md = []
        
        md.append("# 会议汇总对比报告\n")
        md.append(f"生成时间: {datetime.now().isoformat()}  \n")
        md.append(f"会议数量: {len(results)}  \n")
        md.append("\n---\n")
        
        md.append("## 效率指标对比\n")
        md.append("| 会议文件 | 参会人数 | 总发言次数 | 总字数 | 预估时长(分钟) | 平均发言长度 |\n")
        md.append("|----------|----------|------------|--------|----------------|--------------|\n")
        
        for result in results:
            metadata = result.get("metadata", {})
            avg_length = 0
            if metadata.get("total_utterances", 0) > 0:
                avg_length = round(metadata.get("total_words", 0) / metadata.get("total_utterances", 1), 1)
            
            md.append(f"| {metadata.get('filename', 'N/A')} | {metadata.get('speaker_count', 0)} | "
                     f"{metadata.get('total_utterances', 0)} | {metadata.get('total_words', 0)} | "
                     f"{metadata.get('estimated_duration_minutes', 0)} | {avg_length} |\n")
        
        md.append("\n---\n")
        
        md.append("## 参会人员活跃度对比\n")
        
        all_speakers = set()
        for result in results:
            stats = result.get("speaker_statistics", {})
            all_speakers.update(stats.keys())
        
        if all_speakers:
            md.append("| 发言人 | " + " | ".join([r.get("metadata", {}).get("filename", "N/A") for r in results]) + " |\n")
            md.append("|--------|" + "|".join(["---"] * len(results)) + "|\n")
            
            for speaker in sorted(all_speakers):
                row = [speaker]
                for result in results:
                    stats = result.get("speaker_statistics", {})
                    if speaker in stats:
                        count = stats[speaker].get("utterance_count", 0)
                        row.append(f"{count}次")
                    else:
                        row.append("-")
                md.append("| " + " | ".join(row) + " |\n")
        
        md.append("\n---\n")
        
        md.append("## 待办事项汇总\n")
        all_todos = []
        for result in results:
            filename = result.get("metadata", {}).get("filename", "N/A")
            todos = result.get("todo_items", [])
            for todo in todos:
                all_todos.append({
                    "meeting": filename,
                    **todo
                })
        
        if all_todos:
            md.append("| 来源会议 | 任务 | 负责人 | 截止时间 |\n")
            md.append("|----------|------|--------|----------|\n")
            for todo in all_todos:
                md.append(f"| {todo.get('meeting', '')} | {todo.get('task', '')} | "
                         f"{todo.get('assignee', '')} | {todo.get('deadline', '')} |\n")
        else:
            md.append("(无待办事项)\n")
        
        md.append("\n---\n")
        
        md.append("## 关键时间节点汇总\n")
        all_events = []
        for result in results:
            filename = result.get("metadata", {}).get("filename", "N/A")
            events = result.get("time_events", [])
            for event in events:
                all_events.append({
                    "meeting": filename,
                    **event
                })
        
        if all_events:
            md.append("| 来源会议 | 事件 | 时间 |\n")
            md.append("|----------|------|------|\n")
            for event in all_events:
                md.append(f"| {event.get('meeting', '')} | {event.get('event', '')} | {event.get('time', '')} |\n")
        else:
            md.append("(无时间节点)\n")
        
        return "".join(md)
    
    def _generate_comparison_json(self, results: List[Dict[str, Any]]) -> Dict:
        return {
            "report_generated_at": datetime.now().isoformat(),
            "meeting_count": len(results),
            "efficiency_comparison": [
                {
                    "filename": r.get("metadata", {}).get("filename", ""),
                    "speaker_count": r.get("metadata", {}).get("speaker_count", 0),
                    "total_utterances": r.get("metadata", {}).get("total_utterances", 0),
                    "total_words": r.get("metadata", {}).get("total_words", 0),
                    "estimated_duration": r.get("metadata", {}).get("estimated_duration_minutes", 0)
                }
                for r in results
            ],
            "all_todos": [
                {
                    "meeting": r.get("metadata", {}).get("filename", ""),
                    **todo
                }
                for r in results
                for todo in r.get("todo_items", [])
            ],
            "all_time_events": [
                {
                    "meeting": r.get("metadata", {}).get("filename", ""),
                    **event
                }
                for r in results
                for event in r.get("time_events", [])
            ]
        }
