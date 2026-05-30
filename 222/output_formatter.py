import json
import os
from typing import Dict, List, Any
from datetime import datetime

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False


class OutputFormatter:
    @staticmethod
    def to_json(result: Dict[str, Any], indent: int = 2) -> str:
        return json.dumps(result, ensure_ascii=False, indent=indent)
    
    @staticmethod
    def to_markdown(result: Dict[str, Any]) -> str:
        md = []
        
        metadata = result.get("metadata", {})
        md.append(f"# 会议摘要\n")
        md.append(f"**文件**: {metadata.get('filename', 'N/A')}  \n")
        md.append(f"**参会人数**: {metadata.get('speaker_count', 0)}人  \n")
        md.append(f"**总发言次数**: {metadata.get('total_utterances', 0)}次  \n")
        md.append(f"**预估时长**: {metadata.get('estimated_duration_minutes', 0)}分钟  \n")
        md.append(f"**生成时间**: {metadata.get('parsed_at', datetime.now().isoformat())}  \n")
        md.append("\n---\n")
        
        if result.get("overall_summary"):
            md.append("## 会议整体摘要\n")
            md.append(f"{result['overall_summary']}\n")
            md.append("\n---\n")
        
        md.append("## 参会人员统计\n")
        md.append("| 发言人 | 发言次数 | 总字数 | 预估时长(分钟) | 平均字数/次 |\n")
        md.append("|--------|----------|--------|----------------|------------|\n")
        
        stats = result.get("speaker_statistics", {})
        for speaker, stat in stats.items():
            md.append(f"| {speaker} | {stat.get('utterance_count', 0)} | {stat.get('total_words', 0)} | "
                     f"{stat.get('estimated_duration_minutes', 0)} | {stat.get('avg_words_per_utterance', 0)} |\n")
        md.append("\n---\n")
        
        md.append("## 按发言人总结\n")
        summary = result.get("summary", {})
        if summary:
            for speaker, data in summary.items():
                md.append(f"### {speaker}\n")
                
                key_points = data.get("key_points", [])
                if key_points:
                    md.append("**核心观点**:\n")
                    for point in key_points:
                        md.append(f"- {point}\n")
                    md.append("\n")
                
                questions = data.get("questions", [])
                if questions:
                    md.append("**提出的问题**:\n")
                    for q in questions:
                        md.append(f"- {q}\n")
                    md.append("\n")
        else:
            md.append("*(无AI生成的总结数据)*\n")
        
        md.append("\n---\n")
        
        consensus = result.get("meeting_consensus", [])
        if consensus:
            md.append("## 会议达成的共识\n")
            for item in consensus:
                md.append(f"- {item}\n")
            md.append("\n---\n")
        
        todos = result.get("todo_items", [])
        if todos:
            md.append("## 待办事项清单\n")
            md.append("| 序号 | 任务 | 负责人 | 截止时间 |\n")
            md.append("|------|------|--------|----------|\n")
            for idx, todo in enumerate(todos, 1):
                md.append(f"| {idx} | {todo.get('task', '')} | {todo.get('assignee', '')} | {todo.get('deadline', '')} |\n")
            md.append("\n---\n")
        
        time_events = result.get("time_events", [])
        if time_events:
            md.append("## 关键时间节点\n")
            md.append("| 序号 | 事件 | 时间 | 原文 |\n")
            md.append("|------|------|------|------|\n")
            for idx, event in enumerate(time_events, 1):
                md.append(f"| {idx} | {event.get('event', '')} | {event.get('time', '')} | {event.get('original_text', '')} |\n")
            md.append("\n")
        
        return "".join(md)
    
    @staticmethod
    def to_plain_text(result: Dict[str, Any]) -> str:
        text = []
        
        metadata = result.get("metadata", {})
        text.append("=" * 60)
        text.append("会议摘要")
        text.append("=" * 60)
        text.append(f"文件: {metadata.get('filename', 'N/A')}")
        text.append(f"参会人数: {metadata.get('speaker_count', 0)}人")
        text.append(f"总发言次数: {metadata.get('total_utterances', 0)}次")
        text.append(f"预估时长: {metadata.get('estimated_duration_minutes', 0)}分钟")
        text.append(f"生成时间: {metadata.get('parsed_at', datetime.now().isoformat())}")
        text.append("")
        
        if result.get("overall_summary"):
            text.append("-" * 60)
            text.append("会议整体摘要")
            text.append("-" * 60)
            text.append(result['overall_summary'])
            text.append("")
        
        text.append("-" * 60)
        text.append("参会人员统计")
        text.append("-" * 60)
        
        stats = result.get("speaker_statistics", {})
        for speaker, stat in stats.items():
            text.append(f"  {speaker}:")
            text.append(f"    发言次数: {stat.get('utterance_count', 0)}次")
            text.append(f"    总字数: {stat.get('total_words', 0)}字")
            text.append(f"    预估时长: {stat.get('estimated_duration_minutes', 0)}分钟")
            text.append(f"    平均字数/次: {stat.get('avg_words_per_utterance', 0)}")
            text.append("")
        
        text.append("-" * 60)
        text.append("按发言人总结")
        text.append("-" * 60)
        
        summary = result.get("summary", {})
        if summary:
            for speaker, data in summary.items():
                text.append(f"\n【{speaker}】")
                
                key_points = data.get("key_points", [])
                if key_points:
                    text.append("  核心观点:")
                    for point in key_points:
                        text.append(f"    - {point}")
                
                questions = data.get("questions", [])
                if questions:
                    text.append("  提出的问题:")
                    for q in questions:
                        text.append(f"    - {q}")
        else:
            text.append("  (无AI生成的总结数据)")
        
        text.append("")
        text.append("-" * 60)
        
        consensus = result.get("meeting_consensus", [])
        if consensus:
            text.append("会议达成的共识")
            text.append("-" * 60)
            for item in consensus:
                text.append(f"  - {item}")
            text.append("")
        
        todos = result.get("todo_items", [])
        if todos:
            text.append("-" * 60)
            text.append("待办事项清单")
            text.append("-" * 60)
            for idx, todo in enumerate(todos, 1):
                text.append(f"  {idx}. [{todo.get('assignee', '未分配')}] {todo.get('task', '')}")
                if todo.get('deadline'):
                    text.append(f"     截止: {todo.get('deadline', '')}")
            text.append("")
        
        time_events = result.get("time_events", [])
        if time_events:
            text.append("-" * 60)
            text.append("关键时间节点")
            text.append("-" * 60)
            for idx, event in enumerate(time_events, 1):
                text.append(f"  {idx}. [{event.get('time', '')}] {event.get('event', '')}")
        
        text.append("\n" + "=" * 60)
        
        return "\n".join(text)
    
    @staticmethod
    def to_excel(result: Dict[str, Any], output_path: str) -> bool:
        if not OPENPYXL_AVAILABLE:
            raise ImportError("请安装 openpyxl: pip install openpyxl")
        
        wb = Workbook()
        
        header_font = Font(bold=True, size=12, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        center_align = Alignment(horizontal="center", vertical="center")
        wrap_align = Alignment(horizontal="left", vertical="top", wrap_text=True)
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        
        def style_header(ws, row=1):
            for cell in ws[row]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = thin_border
        
        ws1 = wb.active
        ws1.title = "会议概览"
        
        metadata = result.get("metadata", {})
        overview_data = [
            ["项目", "内容"],
            ["文件名", metadata.get('filename', 'N/A')],
            ["参会人数", f"{metadata.get('speaker_count', 0)}人"],
            ["总发言次数", f"{metadata.get('total_utterances', 0)}次"],
            ["总字数", f"{metadata.get('total_words', 0)}字"],
            ["预估时长", f"{metadata.get('estimated_duration_minutes', 0)}分钟"],
            ["生成时间", metadata.get('parsed_at', datetime.now().isoformat())]
        ]
        
        for row_data in overview_data:
            ws1.append(row_data)
        
        style_header(ws1)
        ws1.column_dimensions['A'].width = 15
        ws1.column_dimensions['B'].width = 50
        
        for row in ws1.iter_rows(min_row=1, max_row=ws1.max_row, min_col=1, max_col=2):
            for cell in row:
                cell.border = thin_border
        
        ws2 = wb.create_sheet("参会统计")
        ws2.append(["发言人", "发言次数", "总字数", "预估时长(分钟)", "平均字数/次"])
        
        stats = result.get("speaker_statistics", {})
        for speaker, stat in stats.items():
            ws2.append([
                speaker,
                stat.get('utterance_count', 0),
                stat.get('total_words', 0),
                stat.get('estimated_duration_minutes', 0),
                stat.get('avg_words_per_utterance', 0)
            ])
        
        style_header(ws2)
        for col in ['A', 'B', 'C', 'D', 'E']:
            ws2.column_dimensions[col].width = 18
        for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row, min_col=1, max_col=5):
            for cell in row:
                cell.border = thin_border
                cell.alignment = center_align
        
        ws3 = wb.create_sheet("发言人总结")
        ws3.append(["发言人", "核心观点", "提出的问题"])
        
        summary = result.get("summary", {})
        if summary:
            for speaker, data in summary.items():
                key_points = "; ".join(data.get("key_points", []))
                questions = "; ".join(data.get("questions", []))
                ws3.append([speaker, key_points, questions])
        
        style_header(ws3)
        ws3.column_dimensions['A'].width = 15
        ws3.column_dimensions['B'].width = 60
        ws3.column_dimensions['C'].width = 40
        for row in ws3.iter_rows(min_row=2, max_row=ws3.max_row, min_col=1, max_col=3):
            for cell in row:
                cell.border = thin_border
                cell.alignment = wrap_align
        
        ws4 = wb.create_sheet("会议共识")
        ws4.append(["序号", "共识内容"])
        
        consensus = result.get("meeting_consensus", [])
        for idx, item in enumerate(consensus, 1):
            ws4.append([idx, item])
        
        style_header(ws4)
        ws4.column_dimensions['A'].width = 8
        ws4.column_dimensions['B'].width = 80
        for row in ws4.iter_rows(min_row=1, max_row=ws4.max_row, min_col=1, max_col=2):
            for cell in row:
                cell.border = thin_border
                if cell.row > 1:
                    cell.alignment = wrap_align
        
        ws5 = wb.create_sheet("待办事项")
        ws5.append(["序号", "任务", "负责人", "截止时间"])
        
        todos = result.get("todo_items", [])
        for idx, todo in enumerate(todos, 1):
            ws5.append([
                idx,
                todo.get('task', ''),
                todo.get('assignee', ''),
                todo.get('deadline', '')
            ])
        
        style_header(ws5)
        ws5.column_dimensions['A'].width = 8
        ws5.column_dimensions['B'].width = 60
        ws5.column_dimensions['C'].width = 15
        ws5.column_dimensions['D'].width = 20
        for row in ws5.iter_rows(min_row=1, max_row=ws5.max_row, min_col=1, max_col=4):
            for cell in row:
                cell.border = thin_border
                if cell.row > 1:
                    cell.alignment = wrap_align
        
        ws6 = wb.create_sheet("时间节点")
        ws6.append(["序号", "事件", "时间", "原文"])
        
        time_events = result.get("time_events", [])
        for idx, event in enumerate(time_events, 1):
            ws6.append([
                idx,
                event.get('event', ''),
                event.get('time', ''),
                event.get('original_text', '')
            ])
        
        style_header(ws6)
        ws6.column_dimensions['A'].width = 8
        ws6.column_dimensions['B'].width = 40
        ws6.column_dimensions['C'].width = 20
        ws6.column_dimensions['D'].width = 50
        for row in ws6.iter_rows(min_row=1, max_row=ws6.max_row, min_col=1, max_col=4):
            for cell in row:
                cell.border = thin_border
                if cell.row > 1:
                    cell.alignment = wrap_align
        
        dir_path = os.path.dirname(output_path)
        if dir_path and not os.path.exists(dir_path):
            os.makedirs(dir_path, exist_ok=True)
        
        wb.save(output_path)
        return True
    
    @staticmethod
    def save_to_file(content: str, file_path: str):
        dir_path = os.path.dirname(file_path)
        if dir_path and not os.path.exists(dir_path):
            os.makedirs(dir_path, exist_ok=True)
        with open(file_path, 'w', encoding='utf-8') as f:
            f.write(content)
