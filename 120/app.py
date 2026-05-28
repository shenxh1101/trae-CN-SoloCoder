import os
import json
import uuid
import csv
import io
from datetime import datetime, date
from functools import wraps
from io import BytesIO

import qrcode
from flask import Flask, render_template, request, redirect, url_for, jsonify, send_file, make_response, abort
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

app = Flask(__name__)
app.config['SECRET_KEY'] = 'your-secret-key-here'
app.config['SURVEYS_DIR'] = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'surveys')
app.config['BACKUPS_DIR'] = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'backups')

os.makedirs(app.config['SURVEYS_DIR'], exist_ok=True)
os.makedirs(app.config['BACKUPS_DIR'], exist_ok=True)

TEMPLATES = {
    'satisfaction': {
        'name': '满意度调查',
        'title': '客户满意度调查问卷',
        'description': '感谢您参与我们的满意度调查，您的反馈对我们非常重要！',
        'questions': [
            {
                'type': 'radio',
                'title': '您对我们的服务整体满意度如何？',
                'required': True,
                'options': ['非常满意', '满意', '一般', '不满意', '非常不满意']
            },
            {
                'type': 'checkbox',
                'title': '您认为我们哪些方面需要改进？（可多选）',
                'required': False,
                'options': ['服务态度', '响应速度', '专业能力', '价格合理性', '其他']
            },
            {
                'type': 'text',
                'title': '您有什么其他建议或意见？',
                'required': False
            }
        ]
    },
    'registration': {
        'name': '活动报名',
        'title': '活动报名表',
        'description': '请填写以下信息完成活动报名',
        'questions': [
            {
                'type': 'text',
                'title': '姓名',
                'required': True
            },
            {
                'type': 'radio',
                'title': '性别',
                'required': True,
                'options': ['男', '女', '其他']
            },
            {
                'type': 'text',
                'title': '手机号码',
                'required': True
            },
            {
                'type': 'text',
                'title': '电子邮箱',
                'required': False
            },
            {
                'type': 'checkbox',
                'title': '您感兴趣的活动内容（可多选）',
                'required': False,
                'options': ['主题演讲', '互动讨论', '工作坊', ' networking', '其他']
            }
        ]
    },
    'feedback': {
        'name': '产品反馈',
        'title': '产品使用反馈问卷',
        'description': '请分享您使用我们产品的体验和建议',
        'questions': [
            {
                'type': 'radio',
                'title': '您使用产品的频率是？',
                'required': True,
                'options': ['每天', '每周几次', '每月几次', '偶尔使用', '第一次使用']
            },
            {
                'type': 'radio',
                'title': '您对产品功能的满意度',
                'required': True,
                'options': ['非常满意', '满意', '一般', '不满意', '非常不满意']
            },
            {
                'type': 'text',
                'title': '您最喜欢的功能是什么？',
                'required': False
            },
            {
                'type': 'text',
                'title': '您希望添加什么新功能？',
                'required': False
            }
        ]
    }
}


def get_survey_path(survey_id):
    return os.path.join(app.config['SURVEYS_DIR'], f'{survey_id}.json')


def get_responses_path(survey_id):
    return os.path.join(app.config['SURVEYS_DIR'], f'{survey_id}_responses.json')


def load_survey(survey_id):
    path = get_survey_path(survey_id)
    if os.path.exists(path):
        with open(path, 'r', encoding='utf-8') as f:
            return json.load(f)
    return None


def load_responses(survey_id):
    path = get_responses_path(survey_id)
    if os.path.exists(path):
        with open(path, 'r', encoding='utf-8') as f:
            return json.load(f)
    return []


def save_survey(survey):
    path = get_survey_path(survey['id'])
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(survey, f, ensure_ascii=False, indent=2)


def save_response(survey_id, response):
    path = get_responses_path(survey_id)
    responses = load_responses(survey_id)
    response['timestamp'] = datetime.now().isoformat()
    response['ip'] = request.remote_addr
    responses.append(response)
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(responses, f, ensure_ascii=False, indent=2)
    return len(responses)


def get_client_ip():
    if request.headers.get('X-Forwarded-For'):
        return request.headers.get('X-Forwarded-For').split(',')[0].strip()
    return request.remote_addr


def survey_access_required(f):
    @wraps(f)
    def decorated_function(survey_id, *args, **kwargs):
        survey = load_survey(survey_id)
        if not survey:
            abort(404)
        
        if survey.get('password'):
            if not request.cookies.get(f'auth_{survey_id}'):
                return redirect(url_for('survey_auth', survey_id=survey_id))
        
        if survey.get('expiry_date'):
            try:
                expiry_date = datetime.strptime(survey['expiry_date'], '%Y-%m-%d').date()
                if date.today() > expiry_date:
                    return render_template('expired.html', survey=survey)
            except:
                pass
        
        return f(survey_id, survey, *args, **kwargs)
    return decorated_function


def is_ip_submitted(survey_id):
    responses = load_responses(survey_id)
    ip = get_client_ip()
    return any(r.get('ip') == ip for r in responses)


@app.route('/')
def index():
    surveys = []
    for filename in os.listdir(app.config['SURVEYS_DIR']):
        if filename.endswith('.json') and not filename.endswith('_responses.json'):
            survey_id = filename.replace('.json', '')
            survey = load_survey(survey_id)
            if survey:
                response_count = len(load_responses(survey_id))
                survey['response_count'] = response_count
                surveys.append(survey)
    
    surveys.sort(key=lambda x: x.get('created_at', ''), reverse=True)
    return render_template('index.html', surveys=surveys, templates=TEMPLATES)


@app.route('/create', methods=['GET', 'POST'])
def create_survey():
    if request.method == 'POST':
        data = request.get_json()
        survey_id = str(uuid.uuid4())[:8]
        
        survey = {
            'id': survey_id,
            'title': data['title'],
            'description': data.get('description', ''),
            'questions': data['questions'],
            'created_at': datetime.now().isoformat(),
            'password': data.get('password'),
            'expiry_date': data.get('expiry_date'),
            'one_per_ip': data.get('one_per_ip', False),
            'per_page': data.get('per_page', 0)
        }
        
        save_survey(survey)
        return jsonify({'success': True, 'survey_id': survey_id})
    
    return render_template('create.html', templates=TEMPLATES)


@app.route('/template/<template_id>')
def use_template(template_id):
    if template_id not in TEMPLATES:
        return redirect(url_for('index'))
    template = TEMPLATES[template_id]
    return render_template('create.html', template=template, templates=TEMPLATES)


@app.route('/copy/<survey_id>')
def copy_survey(survey_id):
    survey = load_survey(survey_id)
    if not survey:
        return redirect(url_for('index'))
    
    survey['title'] = survey['title'] + ' (副本)'
    return render_template('create.html', template=survey, templates=TEMPLATES)


@app.route('/survey/<survey_id>/auth', methods=['GET', 'POST'])
def survey_auth(survey_id):
    survey = load_survey(survey_id)
    if not survey:
        abort(404)
    
    if request.method == 'POST':
        password = request.form.get('password', '')
        if password == survey.get('password'):
            resp = make_response(redirect(url_for('view_survey', survey_id=survey_id)))
            resp.set_cookie(f'auth_{survey_id}', 'true', max_age=3600*24)
            return resp
        return render_template('auth.html', survey=survey, error=True)
    
    return render_template('auth.html', survey=survey, error=False)


@app.route('/survey/<survey_id>')
@survey_access_required
def view_survey(survey_id, survey):
    if survey.get('one_per_ip') and is_ip_submitted(survey_id):
        return render_template('already_submitted.html', survey=survey)
    
    per_page = survey.get('per_page', 0)
    total_pages = 1
    current_page = 1
    
    if per_page > 0:
        total_pages = (len(survey['questions']) + per_page - 1) // per_page
        current_page = int(request.args.get('page', 1))
        start = (current_page - 1) * per_page
        end = start + per_page
        questions = survey['questions'][start:end]
    else:
        questions = survey['questions']
    
    return render_template('survey.html', 
                         survey=survey, 
                         questions=questions,
                         current_page=current_page,
                         total_pages=total_pages,
                         per_page=per_page)


@app.route('/survey/<survey_id>/submit', methods=['POST'])
@survey_access_required
def submit_survey(survey_id, survey):
    if survey.get('one_per_ip') and is_ip_submitted(survey_id):
        return jsonify({'success': False, 'error': '您已经提交过此问卷'})
    
    answers = request.get_json().get('answers', {})
    
    errors = []
    for i, q in enumerate(survey['questions']):
        q_key = str(i)
        if q.get('required', False):
            if q_key not in answers or not answers[q_key]:
                errors.append(f'问题 {i+1}: "{q["title"]}" 是必填项')
            elif q['type'] == 'checkbox' and len(answers[q_key]) == 0:
                errors.append(f'问题 {i+1}: "{q["title"]}" 是必填项')
    
    if errors:
        return jsonify({'success': False, 'errors': errors})
    
    count = save_response(survey_id, {'answers': answers})
    return jsonify({'success': True, 'count': count})


@app.route('/survey/<survey_id>/qrcode')
def generate_qrcode(survey_id):
    survey = load_survey(survey_id)
    if not survey:
        abort(404)
    
    url = url_for('view_survey', survey_id=survey_id, _external=True)
    qr = qrcode.QRCode(version=1, box_size=10, border=5)
    qr.add_data(url)
    qr.make(fit=True)
    
    img = qr.make_image(fill_color="black", back_color="white")
    buffer = BytesIO()
    img.save(buffer, 'PNG')
    buffer.seek(0)
    
    return send_file(buffer, mimetype='image/png')


@app.route('/survey/<survey_id>/results')
def view_results(survey_id):
    survey = load_survey(survey_id)
    if not survey:
        abort(404)
    
    responses = load_responses(survey_id)
    stats = calculate_stats(survey, responses)
    
    return render_template('results.html', 
                         survey=survey, 
                         responses=responses, 
                         stats=stats)


def calculate_stats(survey, responses):
    stats = []
    for i, question in enumerate(survey['questions']):
        q_stat = {
            'index': i,
            'title': question['title'],
            'type': question['type'],
            'total_responses': len(responses)
        }
        
        if question['type'] in ['radio', 'checkbox']:
            counts = {opt: 0 for opt in question['options']}
            for response in responses:
                answer = response['answers'].get(str(i))
                if answer:
                    if question['type'] == 'checkbox':
                        for opt in answer:
                            if opt in counts:
                                counts[opt] += 1
                    else:
                        if answer in counts:
                            counts[answer] += 1
            
            q_stat['counts'] = counts
            q_stat['percentages'] = {}
            for opt, count in counts.items():
                q_stat['percentages'][opt] = round(count / len(responses) * 100, 1) if responses else 0
        
        stats.append(q_stat)
    
    return stats


@app.route('/survey/<survey_id>/export/csv')
def export_csv(survey_id):
    survey = load_survey(survey_id)
    if not survey:
        abort(404)
    
    responses = load_responses(survey_id)
    
    output = io.StringIO()
    writer = csv.writer(output)
    
    headers = ['提交时间', 'IP地址']
    for q in survey['questions']:
        headers.append(q['title'])
    writer.writerow(headers)
    
    for r in responses:
        row = [r.get('timestamp', ''), r.get('ip', '')]
        for i, q in enumerate(survey['questions']):
            answer = r['answers'].get(str(i), '')
            if isinstance(answer, list):
                answer = ', '.join(answer)
            row.append(answer)
        writer.writerow(row)
    
    output.seek(0)
    return send_file(
        io.BytesIO(output.getvalue().encode('utf-8-sig')),
        mimetype='text/csv',
        as_attachment=True,
        download_name=f"{survey['title']}.csv"
    )


@app.route('/survey/<survey_id>/export/excel')
def export_excel(survey_id):
    survey = load_survey(survey_id)
    if not survey:
        abort(404)
    
    responses = load_responses(survey_id)
    
    wb = Workbook()
    ws = wb.active
    ws.title = '调查结果'
    
    headers = ['提交时间', 'IP地址']
    for q in survey['questions']:
        headers.append(q['title'])
    
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')
    
    for row_idx, r in enumerate(responses, 2):
        ws.cell(row=row_idx, column=1, value=r.get('timestamp', ''))
        ws.cell(row=row_idx, column=2, value=r.get('ip', ''))
        for i, q in enumerate(survey['questions']):
            answer = r['answers'].get(str(i), '')
            if isinstance(answer, list):
                answer = ', '.join(answer)
            ws.cell(row=row_idx, column=i+3, value=answer)
    
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        ws.column_dimensions[column].width = min(max_length + 2, 50)
    
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    
    return send_file(
        buffer,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f"{survey['title']}.xlsx"
    )


@app.route('/survey/<survey_id>/delete', methods=['POST'])
def delete_survey(survey_id):
    survey = load_survey(survey_id)
    if not survey:
        abort(404)
    
    backup_time = datetime.now().strftime('%Y%m%d_%H%M%S')
    
    survey_path = get_survey_path(survey_id)
    responses_path = get_responses_path(survey_id)
    
    backup_survey = os.path.join(app.config['BACKUPS_DIR'], f'{survey_id}_{backup_time}.json')
    backup_responses = os.path.join(app.config['BACKUPS_DIR'], f'{survey_id}_responses_{backup_time}.json')
    
    with open(survey_path, 'r', encoding='utf-8') as f:
        with open(backup_survey, 'w', encoding='utf-8') as bf:
            bf.write(f.read())
    
    if os.path.exists(responses_path):
        with open(responses_path, 'r', encoding='utf-8') as f:
            with open(backup_responses, 'w', encoding='utf-8') as bf:
                bf.write(f.read())
        os.remove(responses_path)
    
    os.remove(survey_path)
    
    return jsonify({'success': True, 'backup': backup_survey})


@app.route('/survey/<survey_id>/stats')
def get_stats_api(survey_id):
    survey = load_survey(survey_id)
    if not survey:
        abort(404)
    
    responses = load_responses(survey_id)
    stats = calculate_stats(survey, responses)
    
    return jsonify({
        'total': len(responses),
        'stats': stats
    })


@app.errorhandler(404)
def not_found(e):
    return render_template('404.html'), 404


if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5001)
