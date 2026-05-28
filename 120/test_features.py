#!/usr/bin/env python3
"""功能测试脚本 - 验证所有19项功能"""

import json
import sys
import time
import requests
import io
import csv
from openpyxl import load_workbook

BASE_URL = "http://127.0.0.1:5001"

def test_step(step_num, description, func):
    """测试步骤包装器"""
    print(f"\n{'='*60}")
    print(f"测试 {step_num}: {description}")
    print('='*60)
    try:
        result = func()
        print(f"✅ 成功: {description}")
        return result
    except Exception as e:
        print(f"❌ 失败: {description}")
        print(f"   错误: {str(e)}")
        import traceback
        traceback.print_exc()
        return None

def test_1_homepage():
    """测试1: 首页访问"""
    r = requests.get(f"{BASE_URL}/")
    assert r.status_code == 200
    assert "问卷调查系统" in r.text
    assert "创建问卷" in r.text
    assert "快速模板" in r.text
    return True

def test_2_create_page():
    """测试2: 创建问卷页面"""
    r = requests.get(f"{BASE_URL}/create")
    assert r.status_code == 200
    assert "创建问卷" in r.text
    assert "问题列表" in r.text
    return True

def test_3_templates():
    """测试3: 问卷模板功能"""
    for template_id in ['satisfaction', 'registration', 'feedback']:
        r = requests.get(f"{BASE_URL}/template/{template_id}")
        assert r.status_code == 200
        assert "编辑问卷" in r.text
    return True

def test_4_create_survey():
    """测试4: 创建问卷API"""
    survey_data = {
        "title": "功能测试问卷",
        "description": "这是一个自动化测试问卷",
        "password": "",
        "expiry_date": "",
        "per_page": 0,
        "one_per_ip": False,
        "questions": [
            {
                "type": "radio",
                "title": "您喜欢的编程语言？",
                "required": True,
                "options": ["Python", "JavaScript", "Java", "Go"]
            },
            {
                "type": "checkbox",
                "title": "您用过哪些框架？（可多选）",
                "required": False,
                "options": ["Flask", "Django", "React", "Vue", "Spring"]
            },
            {
                "type": "text",
                "title": "您的建议",
                "required": False
            }
        ]
    }
    
    r = requests.post(f"{BASE_URL}/create", 
                      json=survey_data,
                      headers={"Content-Type": "application/json"})
    assert r.status_code == 200
    data = r.json()
    assert data["success"] == True
    assert "survey_id" in data
    print(f"   生成问卷ID: {data['survey_id']}")
    return data["survey_id"]

def test_5_view_survey(survey_id):
    """测试5: 查看问卷页面"""
    r = requests.get(f"{BASE_URL}/survey/{survey_id}")
    assert r.status_code == 200
    assert "功能测试问卷" in r.text
    assert "您喜欢的编程语言" in r.text
    assert "Python" in r.text
    return True

def test_6_qrcode(survey_id):
    """测试6: 二维码生成"""
    r = requests.get(f"{BASE_URL}/survey/{survey_id}/qrcode")
    assert r.status_code == 200
    assert "image/png" in r.headers["Content-Type"]
    print(f"   Content-Type: {r.headers['Content-Type']}")
    print(f"   二维码大小: {len(r.content)} bytes")
    assert len(r.content) > 100  # PNG文件应该大于100bytes
    return True

def test_7_submit_survey(survey_id):
    """测试7: 提交问卷答案"""
    answers = {
        "0": "Python",
        "1": ["Flask", "Django"],
        "2": "很好用的系统！"
    }
    
    r = requests.post(f"{BASE_URL}/survey/{survey_id}/submit",
                      json={"answers": answers},
                      headers={"Content-Type": "application/json"})
    assert r.status_code == 200
    data = r.json()
    assert data["success"] == True
    print(f"   提交后总数: {data['count']}")
    return True

def test_8_required_validation(survey_id):
    """测试8: 必填验证"""
    answers = {
        "1": ["React"],
        "2": "跳过必填项"
    }
    
    r = requests.post(f"{BASE_URL}/survey/{survey_id}/submit",
                      json={"answers": answers},
                      headers={"Content-Type": "application/json"})
    assert r.status_code == 200
    data = r.json()
    assert data["success"] == False
    assert "errors" in data or "error" in data
    print(f"   验证失败提示: {data.get('errors', data.get('error'))}")
    return True

def test_9_view_results(survey_id):
    """测试9: 查看结果页面"""
    r = requests.get(f"{BASE_URL}/survey/{survey_id}/results")
    assert r.status_code == 200
    assert "调查结果" in r.text
    assert "统计图表" in r.text
    assert "详细回答" in r.text
    return True

def test_10_stats_api(survey_id):
    """测试10: 统计数据API"""
    r = requests.get(f"{BASE_URL}/survey/{survey_id}/stats")
    assert r.status_code == 200
    data = r.json()
    assert "total" in data
    assert "stats" in data
    assert data["total"] >= 1
    print(f"   总回答数: {data['total']}")
    print(f"   统计问题数: {len(data['stats'])}")
    return data

def test_11_radio_stats(stats_data):
    """测试11: 单选题统计"""
    radio_stat = stats_data["stats"][0]
    assert radio_stat["type"] == "radio"
    assert "counts" in radio_stat
    assert "percentages" in radio_stat
    assert "Python" in radio_stat["counts"]
    print(f"   Python选择数: {radio_stat['counts']['Python']}")
    print(f"   Python百分比: {radio_stat['percentages']['Python']}%")
    return True

def test_12_checkbox_stats(stats_data):
    """测试12: 多选题统计"""
    checkbox_stat = stats_data["stats"][1]
    assert checkbox_stat["type"] == "checkbox"
    assert "counts" in checkbox_stat
    assert "percentages" in checkbox_stat
    print(f"   Flask选择数: {checkbox_stat['counts']['Flask']}")
    return True

def test_13_export_csv(survey_id):
    """测试13: CSV导出"""
    r = requests.get(f"{BASE_URL}/survey/{survey_id}/export/csv")
    assert r.status_code == 200
    print(f"   Content-Type: {r.headers['Content-Type']}")
    assert "text/csv" in r.headers["Content-Type"]
    content = r.content.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(content))
    rows = list(reader)
    assert len(rows) >= 2  # 表头 + 至少一行数据
    assert "提交时间" in rows[0]
    assert "IP地址" in rows[0]
    print(f"   CSV行数: {len(rows)}")
    return True

def test_14_export_excel(survey_id):
    """测试14: Excel导出"""
    r = requests.get(f"{BASE_URL}/survey/{survey_id}/export/excel")
    assert r.status_code == 200
    assert "openxmlformats-officedocument" in r.headers["Content-Type"]
    
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    assert ws.title == "调查结果"
    
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) >= 2
    assert "提交时间" in rows[0]
    assert "IP地址" in rows[0]
    print(f"   Excel行数: {len(rows)}")
    print(f"   Excel列数: {len(rows[0])}")
    return True

def test_15_password_protected_survey():
    """测试15: 密码保护问卷"""
    survey_data = {
        "title": "密码保护测试问卷",
        "description": "需要密码才能访问",
        "password": "test123",
        "expiry_date": "",
        "per_page": 0,
        "one_per_ip": False,
        "questions": [
            {"type": "radio", "title": "测试问题", "required": True, "options": ["A", "B"]}
        ]
    }
    
    r = requests.post(f"{BASE_URL}/create", json=survey_data,
                      headers={"Content-Type": "application/json"})
    survey_id = r.json()["survey_id"]
    
    r2 = requests.get(f"{BASE_URL}/survey/{survey_id}", allow_redirects=False)
    assert r2.status_code == 302  # 重定向到密码验证页面
    assert "/auth" in r2.headers["Location"]
    print(f"   密码问卷ID: {survey_id}")
    return survey_id

def test_16_expiry_date():
    """测试16: 截止日期功能"""
    survey_data = {
        "title": "过期问卷测试",
        "description": "已过期的问卷",
        "password": "",
        "expiry_date": "2020-01-01",  # 过去的日期
        "per_page": 0,
        "one_per_ip": False,
        "questions": [
            {"type": "radio", "title": "测试问题", "required": True, "options": ["A", "B"]}
        ]
    }
    
    r = requests.post(f"{BASE_URL}/create", json=survey_data,
                      headers={"Content-Type": "application/json"})
    survey_id = r.json()["survey_id"]
    
    r2 = requests.get(f"{BASE_URL}/survey/{survey_id}")
    assert r2.status_code == 200
    assert "问卷已过期" in r2.text
    print(f"   过期问卷ID: {survey_id}")
    return survey_id

def test_17_pagination():
    """测试17: 问卷分页功能"""
    questions = []
    for i in range(5):
        questions.append({
            "type": "radio",
            "title": f"分页测试问题{i+1}",
            "required": True,
            "options": ["A", "B", "C"]
        })
    
    survey_data = {
        "title": "分页测试问卷",
        "description": "每页显示2个问题",
        "password": "",
        "expiry_date": "",
        "per_page": 2,
        "one_per_ip": False,
        "questions": questions
    }
    
    r = requests.post(f"{BASE_URL}/create", json=survey_data,
                      headers={"Content-Type": "application/json"})
    survey_id = r.json()["survey_id"]
    
    r1 = requests.get(f"{BASE_URL}/survey/{survey_id}?page=1")
    assert "分页测试问题1" in r1.text
    assert "分页测试问题2" in r1.text
    assert "分页测试问题3" not in r1.text
    
    r2 = requests.get(f"{BASE_URL}/survey/{survey_id}?page=2")
    assert "分页测试问题3" in r2.text
    assert "分页测试问题4" in r2.text
    
    print(f"   分页问卷ID: {survey_id}")
    return survey_id

def test_18_copy_survey(survey_id):
    """测试18: 复制问卷功能"""
    r = requests.get(f"{BASE_URL}/copy/{survey_id}")
    assert r.status_code == 200
    assert "编辑问卷" in r.text
    assert "副本" in r.text or "功能测试问卷" in r.text
    return True

def test_19_delete_survey(survey_id):
    """测试19: 删除问卷功能（含备份）"""
    r = requests.post(f"{BASE_URL}/survey/{survey_id}/delete")
    assert r.status_code == 200
    data = r.json()
    assert data["success"] == True
    assert "backup" in data
    print(f"   备份文件: {data['backup']}")
    
    r2 = requests.get(f"{BASE_URL}/survey/{survey_id}")
    assert r2.status_code == 404  # 问卷已删除
    return True

def main():
    print("""
╔══════════════════════════════════════════════════════════════╗
║              问卷调查系统 - 完整功能测试套件                  ║
╚══════════════════════════════════════════════════════════════╝""")
    
    # 等待服务启动
    time.sleep(1)
    
    passed = 0
    total = 19
    
    # 1-3: 基础功能
    test_step(1, "首页访问", test_1_homepage) and (passed := passed + 1)
    test_step(2, "创建问卷页面", test_2_create_page) and (passed := passed + 1)
    test_step(3, "问卷模板功能", test_3_templates) and (passed := passed + 1)
    
    # 4-8: 问卷创建和提交
    survey_id = test_step(4, "创建问卷API", test_4_create_survey)
    if survey_id: passed += 1
    
    test_step(5, "查看问卷页面", lambda: test_5_view_survey(survey_id)) and (passed := passed + 1)
    test_step(6, "二维码生成", lambda: test_6_qrcode(survey_id)) and (passed := passed + 1)
    test_step(7, "提交问卷答案", lambda: test_7_submit_survey(survey_id)) and (passed := passed + 1)
    test_step(8, "必填验证功能", lambda: test_8_required_validation(survey_id)) and (passed := passed + 1)
    
    # 9-14: 结果和导出
    test_step(9, "查看结果页面", lambda: test_9_view_results(survey_id)) and (passed := passed + 1)
    stats_data = test_step(10, "统计数据API", lambda: test_10_stats_api(survey_id))
    if stats_data: passed += 1
    
    test_step(11, "单选题统计", lambda: test_11_radio_stats(stats_data)) and (passed := passed + 1)
    test_step(12, "多选题统计", lambda: test_12_checkbox_stats(stats_data)) and (passed := passed + 1)
    test_step(13, "CSV导出功能", lambda: test_13_export_csv(survey_id)) and (passed := passed + 1)
    test_step(14, "Excel导出功能", lambda: test_14_export_excel(survey_id)) and (passed := passed + 1)
    
    # 15-17: 高级功能
    pwd_survey_id = test_step(15, "密码保护功能", test_15_password_protected_survey)
    if pwd_survey_id: passed += 1
    
    expiry_survey_id = test_step(16, "截止日期功能", test_16_expiry_date)
    if expiry_survey_id: passed += 1
    
    page_survey_id = test_step(17, "问卷分页功能", test_17_pagination)
    if page_survey_id: passed += 1
    
    # 18-19: 管理功能
    test_step(18, "复制问卷功能", lambda: test_18_copy_survey(survey_id)) and (passed := passed + 1)
    test_step(19, "删除问卷功能", lambda: test_19_delete_survey(survey_id)) and (passed := passed + 1)
    
    # 总结
    print(f"\n{'='*60}")
    print(f"测试完成: {passed}/{total} 项功能通过")
    print('='*60)
    
    if passed == total:
        print("\n🎉 所有功能测试通过！系统完整可用。")
        return 0
    else:
        print(f"\n⚠️  {total - passed} 项功能需要检查。")
        return 1

if __name__ == "__main__":
    sys.exit(main())
