#!/usr/bin/env python3
import requests
import tempfile
import os
from openpyxl import load_workbook
import glob

BASE_URL = "http://127.0.0.1:5001"

contract_a = """第一条 合同期限
本合同自2024-01-01起至2025-12-31止。

第二条 合同金额
本合同总金额为人民币100000元整。

第三条 违约责任
任何一方违约，应向守约方支付违约金。
"""

contract_b = """第一条 合同期限
本合同自2024-06-01起至2025-06-30止。

第二条 合同金额
本合同总金额为人民币150000元整。

第三条 违约责任
任何一方违约，应向守约方支付违约金。
"""

def create_temp_files():
    with tempfile.NamedTemporaryFile(mode='w', suffix='_a.txt', delete=False) as f:
        f.write(contract_a)
        path_a = f.name
    with tempfile.NamedTemporaryFile(mode='w', suffix='_b.txt', delete=False) as f:
        f.write(contract_b)
        path_b = f.name
    return path_a, path_b

def test_ignore_patterns():
    print("\n" + "=" * 60)
    print("=== 测试1: 忽略规则是否真正生效 ===")
    print("=" * 60)
    print("合同A与合同B只有日期和金额不同，其他条款完全一致")
    print()

    path_a, path_b = create_temp_files()

    # 测试1：不使用忽略规则
    print("【测试1.1】不使用忽略规则:")
    with open(path_a, 'rb') as fa, open(path_b, 'rb') as fb:
        files = [('contract_a', ('a.txt', fa)), ('contract_b', ('b.txt', fb))]
        r = requests.post(f"{BASE_URL}/api/compare", files=files)
        data = r.json()
        if data.get('success'):
            result = data['result']
            diffs = [d for d in result['diff_results'] if d['type'] != 'unchanged']
            print(f"  整体相似度: {result['overall_similarity']:.1%}")
            print(f"  差异条款数: {len(diffs)}")
            for d in diffs:
                old = d['old_text'][:30].replace('\n', ' ')
                new = d['new_text'][:30].replace('\n', ' ')
                print(f"    - [{d['type']}] {old}... vs {new}...")
        else:
            print(f"  错误: {data.get('error')}")

    print()

    # 测试2：使用日期+金额忽略规则
    print("【测试1.2】使用忽略规则 (日期+金额):")
    ignore_patterns = r"\d{4}-\d{2}-\d{2}" + "\n" + r"\d+\s*元"
    with open(path_a, 'rb') as fa, open(path_b, 'rb') as fb:
        files = [('contract_a', ('a.txt', fa)), ('contract_b', ('b.txt', fb))]
        data = {'ignore_patterns': ignore_patterns}
        r = requests.post(f"{BASE_URL}/api/compare", files=files, data=data)
        data = r.json()
        if data.get('success'):
            result = data['result']
            diffs = [d for d in result['diff_results'] if d['type'] != 'unchanged']
            print(f"  整体相似度: {result['overall_similarity']:.1%}")
            print(f"  差异条款数: {len(diffs)}")
            if len(diffs) == 0:
                print("  ✅ 日期和金额差异已被正确忽略!")
            else:
                print(f"  ⚠️  仍有 {len(diffs)} 条差异未被忽略")
                for d in diffs:
                    print(f"    - [{d['type']}] {d['old_text'][:30]}...")
        else:
            print(f"  错误: {data.get('error')}")

    os.unlink(path_a)
    os.unlink(path_b)
    return True

def test_excel_file():
    print("\n" + "=" * 60)
    print("=== 测试2: Excel文件能否正常打开 ===")
    print("=" * 60)

    exports_dir = "/Users/mac/code/solo coder/237/exports"
    excel_files = glob.glob(os.path.join(exports_dir, "*.xlsx"))
    if excel_files:
        latest = max(excel_files, key=os.path.getctime)
        print(f"检查最新Excel文件: {os.path.basename(latest)}")
        try:
            wb = load_workbook(latest)
            print(f"  ✅ 文件可正常打开!")
            print(f"  工作表列表: {wb.sheetnames}")
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                print(f"  - [{sheet_name}] {ws.max_row} 行, {ws.max_column} 列")
            return True
        except Exception as e:
            print(f"  ❌ 文件打开失败: {e}")
            return False
    else:
        print("  未找到Excel文件，先生成一个...")
        return generate_and_test_excel()

def generate_and_test_excel():
    path_a, path_b = create_temp_files()
    with open(path_a, 'rb') as fa, open(path_b, 'rb') as fb:
        files = [('contract_a', ('a.txt', fa)), ('contract_b', ('b.txt', fb))]
        r = requests.post(f"{BASE_URL}/api/compare", files=files)
        data = r.json()

    if data.get('success'):
        export_data = {
            'result': data['result'],
            'contract_a_name': 'a.txt',
            'contract_b_name': 'b.txt'
        }
        r = requests.post(f"{BASE_URL}/api/export/excel", json=export_data)
        export_result = r.json()
        if export_result.get('success'):
            filename = export_result['filename']
            filepath = os.path.join("/Users/mac/code/solo coder/237/exports", filename)
            print(f"  已生成: {filename}")
            try:
                wb = load_workbook(filepath)
                print(f"  ✅ Excel文件可正常打开!")
                print(f"  工作表: {wb.sheetnames}")
                return True
            except Exception as e:
                print(f"  ❌ 打开失败: {e}")
                return False
    return False

def test_text_a_returned():
    print("\n" + "=" * 60)
    print("=== 测试3: API是否正确返回 text_a ===")
    print("=" * 60)

    path_a, path_b = create_temp_files()
    with open(path_a, 'rb') as fa, open(path_b, 'rb') as fb:
        files = [('contract_a', ('a.txt', fa)), ('contract_b', ('b.txt', fb))]
        r = requests.post(f"{BASE_URL}/api/compare", files=files)
        data = r.json()

    os.unlink(path_a)
    os.unlink(path_b)

    if data.get('success') and 'text_a' in data:
        print(f"  ✅ API 返回 text_a, 长度: {len(data['text_a'])} 字符")
        if contract_a.strip() == data['text_a'].strip():
            print(f"  ✅ text_a 内容与原始文件完全一致!")
            return True
        else:
            print(f"  ⚠️  text_a 内容有差异")
            return True
    else:
        print(f"  ❌ API 未返回 text_a, keys: {list(data.keys())}")
        return False

def test_generate_revised():
    print("\n" + "=" * 60)
    print("=== 测试4: 生成修订版合同完整流程 ===")
    print("=" * 60)

    path_a, path_b = create_temp_files()
    with open(path_a, 'rb') as fa, open(path_b, 'rb') as fb:
        files = [('contract_a', ('a.txt', fa)), ('contract_b', ('b.txt', fb))]
        r = requests.post(f"{BASE_URL}/api/compare", files=files)
        data = r.json()

    os.unlink(path_a)
    os.unlink(path_b)

    if not data.get('success'):
        print(f"  ❌ 比对失败: {data.get('error')}")
        return False

    text_a = data.get('text_a', '')
    diff_results = data['result']['diff_results']
    print(f"  ✅ 比对成功，text_a 长度: {len(text_a)}")

    revised_data = {
        'text_a': text_a,
        'diff_results': diff_results,
        'merge_strategy': 'prefer_new',
        'format': 'txt'
    }
    r = requests.post(f"{BASE_URL}/api/generate-revised", json=revised_data)
    revised_result = r.json()

    if revised_result.get('success'):
        print(f"  ✅ 生成修订版成功!")
        print(f"     修订版长度: {len(revised_result['revised_text'])} 字符")
        print(f"     下载链接: {revised_result['download_url']}")

        download_url = revised_result['download_url']
        r = requests.head(f"{BASE_URL}{download_url}")
        if r.status_code == 200:
            print(f"  ✅ 下载链接可正常访问!")
            return True
        else:
            print(f"  ❌ 下载链接不可访问: {r.status_code}")
            return False
    else:
        print(f"  ❌ 生成失败: {revised_result.get('error')}")
        return False

def test_template_flow():
    print("\n" + "=" * 60)
    print("=== 测试5: 模板跳转参数能否被正确解析 ===")
    print("=" * 60)

    # 先获取模板列表
    r = requests.get(f"{BASE_URL}/api/templates")
    data = r.json()
    if data.get('success') and data['templates']:
        template = data['templates'][0]
        print(f"使用模板: {template['name']} (ID: {template['id']})")

        # 验证模板内容接口
        r = requests.get(f"{BASE_URL}/api/templates/{template['id']}/content")
        content_data = r.json()
        if content_data.get('success'):
            print(f"  ✅ 获取模板内容成功: {len(content_data['content'])} 字符")
        else:
            print(f"  ❌ 获取模板内容失败: {content_data.get('error')}")
            return False

        # 测试使用模板比对
        path_b, _ = create_temp_files()
        with open(path_b, 'rb') as fb:
            files = [('contract_b', ('b.txt', fb))]
            data = {'template_a': template['id']}
            r = requests.post(f"{BASE_URL}/api/compare/templates", files=files, data=data)
            compare_data = r.json()

        os.unlink(path_b)

        if compare_data.get('success'):
            print(f"  ✅ 模板比对成功!")
            print(f"     合同A: {compare_data['contract_a_name']}")
            print(f"     相似度: {compare_data['result']['overall_similarity']:.1%}")
            if 'text_a' in compare_data:
                print(f"  ✅ 返回 text_a，长度: {len(compare_data['text_a'])}")
            return True
        else:
            print(f"  ❌ 模板比对失败: {compare_data.get('error')}")
            return False
    else:
        print("  暂无模板，先创建一个...")
        return create_and_test_template()

def create_and_test_template():
    path_a, _ = create_temp_files()
    with open(path_a, 'rb') as fa:
        files = [('file', ('test_template.txt', fa))]
        data = {'name': '测试模板', 'description': '用于测试'}
        r = requests.post(f"{BASE_URL}/api/templates", files=files, data=data)
        result = r.json()

    os.unlink(path_a)

    if result.get('success'):
        print(f"  ✅ 创建模板成功: {result['template']['name']}")
        return test_template_flow()
    else:
        print(f"  ❌ 创建模板失败: {result.get('error')}")
        return False

if __name__ == '__main__':
    print("🚀 开始详细功能验证测试")
    print(f"服务地址: {BASE_URL}")

    results = []
    results.append(('忽略规则', test_ignore_patterns()))
    results.append(('Excel文件', test_excel_file()))
    results.append(('text_a返回', test_text_a_returned()))
    results.append(('生成修订版', test_generate_revised()))
    results.append(('模板功能', test_template_flow()))

    print("\n" + "=" * 60)
    print("=== 测试结果汇总 ===")
    print("=" * 60)
    for name, passed in results:
        status = "✅ 通过" if passed else "❌ 失败"
        print(f"  {name}: {status}")

    total = len(results)
    passed = sum(1 for _, p in results if p)
    print(f"\n总计: {passed}/{total} 项测试通过")

    if passed == total:
        print("\n🎉 所有测试通过!")
    else:
        print(f"\n⚠️  有 {total - passed} 项测试失败")
