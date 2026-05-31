import os
import json
from datetime import datetime

try:
    import pandas as pd
    HAS_PANDAS = True
except ImportError:
    HAS_PANDAS = False

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

from docx import Document

class Exporter:
    def __init__(self, export_folder):
        self.export_folder = export_folder
        os.makedirs(self.export_folder, exist_ok=True)
    
    def generate_filename(self, prefix, ext):
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        return f"{prefix}_{timestamp}.{ext}"
    
    def export_to_html(self, comparison_result, contract_a_name, contract_b_name, focus_categories=None):
        filename = self.generate_filename('comparison_report', 'html')
        filepath = os.path.join(self.export_folder, filename)
        
        html_content = self._generate_html_report(
            comparison_result, contract_a_name, contract_b_name, focus_categories
        )
        
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(html_content)
        
        return filepath, filename
    
    def export_to_excel(self, comparison_result, contract_a_name, contract_b_name):
        filename = self.generate_filename('comparison_report', 'xlsx')
        filepath = os.path.join(self.export_folder, filename)
        
        if HAS_OPENPYXL:
            self._export_to_excel_openpyxl(comparison_result, contract_a_name, contract_b_name, filepath)
        elif HAS_PANDAS:
            self._export_to_excel_pandas(comparison_result, contract_a_name, contract_b_name, filepath)
        else:
            return self.export_to_json(comparison_result, contract_a_name, contract_b_name)
        
        return filepath, filename
    
    def _export_to_excel_openpyxl(self, comparison_result, contract_a_name, contract_b_name, filepath):
        wb = Workbook()
        
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        wrap_align = Alignment(wrap_text=True, vertical="top")
        
        ws_summary = wb.active
        ws_summary.title = "汇总"
        summary_headers = ['指标', '数值']
        for col, h in enumerate(summary_headers, 1):
            cell = ws_summary.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border
        
        summary_rows = [
            ('合同A名称', contract_a_name),
            ('合同B名称', contract_b_name),
            ('合同A条款数', comparison_result['total_clauses_a']),
            ('合同B条款数', comparison_result['total_clauses_b']),
            ('整体相似度', f"{comparison_result['overall_similarity']:.2%}"),
            ('新增条款', comparison_result['stats']['added']),
            ('删除条款', comparison_result['stats']['deleted']),
            ('修改条款', comparison_result['stats']['modified']),
            ('未变化条款', comparison_result['stats']['unchanged']),
        ]
        for row_idx, (label, value) in enumerate(summary_rows, 2):
            ws_summary.cell(row=row_idx, column=1, value=label).border = thin_border
            ws_summary.cell(row=row_idx, column=2, value=value).border = thin_border
        ws_summary.column_dimensions['A'].width = 18
        ws_summary.column_dimensions['B'].width = 20
        
        ws_detail = wb.create_sheet("差异详情")
        detail_headers = ['序号', '类别', '类型', '相似度', '合同A内容', '合同B内容']
        for col, h in enumerate(detail_headers, 1):
            cell = ws_detail.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border
        
        type_fills = {
            'added': PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid"),
            'deleted': PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid"),
            'modified': PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"),
        }
        
        for row_idx, item in enumerate(comparison_result['diff_results'], 2):
            row_data = [
                row_idx - 1,
                item['category'],
                self._get_type_label(item['type']),
                f"{item['similarity']:.2%}",
                item['old_text'],
                item['new_text']
            ]
            fill = type_fills.get(item['type'])
            for col, val in enumerate(row_data, 1):
                cell = ws_detail.cell(row=row_idx, column=col, value=val)
                cell.border = thin_border
                cell.alignment = wrap_align
                if fill:
                    cell.fill = fill
        
        ws_detail.column_dimensions['A'].width = 6
        ws_detail.column_dimensions['B'].width = 12
        ws_detail.column_dimensions['C'].width = 8
        ws_detail.column_dimensions['D'].width = 10
        ws_detail.column_dimensions['E'].width = 40
        ws_detail.column_dimensions['F'].width = 40
        
        wb.save(filepath)
    
    def _export_to_excel_pandas(self, comparison_result, contract_a_name, contract_b_name, filepath):
        data = []
        for idx, item in enumerate(comparison_result['diff_results']):
            data.append({
                '序号': idx + 1,
                '类别': item['category'],
                '类型': self._get_type_label(item['type']),
                '相似度': f"{item['similarity']:.2%}",
                '合同A内容': item['old_text'],
                '合同B内容': item['new_text']
            })
        
        df = pd.DataFrame(data)
        
        summary_data = {
            '指标': ['合同A条款数', '合同B条款数', '整体相似度', '新增条款', '删除条款', '修改条款', '未变化条款'],
            '数值': [
                comparison_result['total_clauses_a'],
                comparison_result['total_clauses_b'],
                f"{comparison_result['overall_similarity']:.2%}",
                comparison_result['stats']['added'],
                comparison_result['stats']['deleted'],
                comparison_result['stats']['modified'],
                comparison_result['stats']['unchanged']
            ]
        }
        df_summary = pd.DataFrame(summary_data)
        
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            df_summary.to_excel(writer, sheet_name='汇总', index=False)
            df.to_excel(writer, sheet_name='差异详情', index=False)
    
    def export_to_json(self, comparison_result, contract_a_name, contract_b_name):
        filename = self.generate_filename('comparison_report', 'json')
        filepath = os.path.join(self.export_folder, filename)
        
        json_data = {
            'contract_a': contract_a_name,
            'contract_b': contract_b_name,
            'generated_at': datetime.now().isoformat(),
            'overall_similarity': comparison_result['overall_similarity'],
            'total_clauses_a': comparison_result['total_clauses_a'],
            'total_clauses_b': comparison_result['total_clauses_b'],
            'stats': comparison_result['stats'],
            'differences': []
        }
        
        for item in comparison_result['diff_results']:
            json_data['differences'].append({
                'category': item['category'],
                'type': item['type'],
                'similarity': item['similarity'],
                'old_text': item['old_text'],
                'new_text': item['new_text']
            })
        
        with open(filepath, 'w', encoding='utf-8') as f:
            json.dump(json_data, f, ensure_ascii=False, indent=2)
        
        return filepath, filename
    
    def export_revised_contract(self, revised_text, format='txt'):
        filename = self.generate_filename('revised_contract', format)
        filepath = os.path.join(self.export_folder, filename)
        
        if format == 'txt':
            with open(filepath, 'w', encoding='utf-8') as f:
                f.write(revised_text)
        elif format == 'docx':
            doc = Document()
            for paragraph in revised_text.split('\n\n'):
                doc.add_paragraph(paragraph)
            doc.save(filepath)
        
        return filepath, filename
    
    def _get_type_label(self, diff_type):
        labels = {
            'added': '新增',
            'deleted': '删除',
            'modified': '修改',
            'unchanged': '未变化'
        }
        return labels.get(diff_type, diff_type)
    
    def _generate_html_report(self, comparison_result, contract_a_name, contract_b_name, focus_categories):
        css = """
        <style>
            body { font-family: Arial, sans-serif; max-width: 1200px; margin: 0 auto; padding: 20px; }
            h1 { color: #333; border-bottom: 2px solid #007bff; padding-bottom: 10px; }
            .summary { background: #f8f9fa; padding: 20px; border-radius: 8px; margin-bottom: 20px; }
            .summary-grid { display: grid; grid-template-columns: repeat(4, 1fr); gap: 15px; margin-top: 15px; }
            .summary-item { text-align: center; padding: 15px; background: white; border-radius: 8px; }
            .summary-item .label { font-size: 14px; color: #666; }
            .summary-item .value { font-size: 24px; font-weight: bold; color: #007bff; }
            .added .value { color: #28a745; }
            .deleted .value { color: #dc3545; }
            .modified .value { color: #ffc107; }
            .similarity-meter { height: 30px; background: #e9ecef; border-radius: 15px; overflow: hidden; margin-top: 10px; }
            .similarity-fill { height: 100%; background: linear-gradient(90deg, #dc3545, #ffc107, #28a745); }
            .diff-item { border: 1px solid #ddd; border-radius: 8px; margin-bottom: 15px; overflow: hidden; }
            .diff-header { padding: 10px 15px; background: #f8f9fa; display: flex; justify-content: space-between; align-items: center; }
            .diff-type { padding: 4px 12px; border-radius: 12px; font-size: 12px; font-weight: bold; color: white; }
            .type-added { background: #28a745; }
            .type-deleted { background: #dc3545; }
            .type-modified { background: #ffc107; color: #333; }
            .type-unchanged { background: #6c757d; }
            .category-tag { padding: 4px 8px; background: #007bff; color: white; border-radius: 4px; font-size: 12px; margin-left: 10px; }
            .diff-content { display: grid; grid-template-columns: 1fr 1fr; gap: 15px; padding: 15px; }
            .diff-pane { padding: 15px; background: #f8f9fa; border-radius: 4px; }
            .diff-pane h4 { margin-top: 0; color: #555; }
            .diff-equal { background: transparent; }
            .diff-insert { background: #d4edda; color: #155724; padding: 2px 4px; border-radius: 3px; }
            .diff-delete { background: #f8d7da; color: #721c24; padding: 2px 4px; border-radius: 3px; text-decoration: line-through; }
            .focus-highlight { border-left: 4px solid #007bff; }
            .legend { display: flex; gap: 20px; margin-top: 10px; }
            .legend-item { display: flex; align-items: center; gap: 5px; }
            .legend-color { width: 20px; height: 20px; border-radius: 3px; }
        </style>
        """
        
        stats = comparison_result['stats']
        similarity_pct = comparison_result['overall_similarity'] * 100
        
        html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>合同比对报告</title>
    {css}
</head>
<body>
    <h1>📄 合同条款智能比对报告</h1>
    
    <div class="summary">
        <h3>比对概览</h3>
        <p><strong>合同A:</strong> {contract_a_name}</p>
        <p><strong>合同B:</strong> {contract_b_name}</p>
        <p><strong>生成时间:</strong> {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>
        
        <div style="margin-top: 15px;">
            <strong>整体相似度: {similarity_pct:.1f}%</strong>
            <div class="similarity-meter">
                <div class="similarity-fill" style="width: {similarity_pct}%"></div>
            </div>
        </div>
        
        <div class="summary-grid">
            <div class="summary-item">
                <div class="label">合同A条款数</div>
                <div class="value">{comparison_result['total_clauses_a']}</div>
            </div>
            <div class="summary-item">
                <div class="label">合同B条款数</div>
                <div class="value">{comparison_result['total_clauses_b']}</div>
            </div>
            <div class="summary-item added">
                <div class="label">新增条款</div>
                <div class="value">{stats['added']}</div>
            </div>
            <div class="summary-item deleted">
                <div class="label">删除条款</div>
                <div class="value">{stats['deleted']}</div>
            </div>
            <div class="summary-item modified">
                <div class="label">修改条款</div>
                <div class="value">{stats['modified']}</div>
            </div>
            <div class="summary-item">
                <div class="label">未变化条款</div>
                <div class="value">{stats['unchanged']}</div>
            </div>
        </div>
        
        <div class="legend">
            <div class="legend-item"><div class="legend-color" style="background: #d4edda;"></div>新增内容</div>
            <div class="legend-item"><div class="legend-color" style="background: #f8d7da;"></div>删除内容</div>
        </div>
        
        {f'<p><strong>关注类别:</strong> {", ".join(focus_categories)}</p>' if focus_categories else ''}
    </div>
    
    <h2>差异详情</h2>
"""
        
        for item in comparison_result['diff_results']:
            is_focus = focus_categories and item['category'] in focus_categories
            focus_class = 'focus-highlight' if is_focus else ''
            
            html += f"""
    <div class="diff-item {focus_class}">
        <div class="diff-header">
            <div>
                <span class="diff-type type-{item['type']}">{self._get_type_label(item['type'])}</span>
                <span class="category-tag">{item['category']}</span>
                {f'<span class="category-tag" style="background: #007bff;">⭐ 重点关注</span>' if is_focus else ''}
            </div>
            <div>相似度: {item['similarity']:.1%}</div>
        </div>
        <div class="diff-content">
            <div class="diff-pane">
                <h4>合同A (原始)</h4>
                <p>{item['old_highlighted'] or '-'}</p>
            </div>
            <div class="diff-pane">
                <h4>合同B (对比)</h4>
                <p>{item['new_highlighted'] or '-'}</p>
            </div>
        </div>
    </div>
"""
        
        html += """
</body>
</html>"""
        
        return html
    
    def export_batch_summary(self, batch_results, base_contract_name):
        filename = self.generate_filename('batch_comparison_summary', 'xlsx')
        filepath = os.path.join(self.export_folder, filename)
        
        if HAS_OPENPYXL:
            self._export_batch_summary_openpyxl(batch_results, base_contract_name, filepath)
        elif HAS_PANDAS:
            self._export_batch_summary_pandas(batch_results, base_contract_name, filepath)
        else:
            return self._export_batch_summary_csv(batch_results, base_contract_name, filepath)
        
        return filepath, filename
    
    def _export_batch_summary_openpyxl(self, batch_results, base_contract_name, filepath):
        wb = Workbook()
        
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center")
        thin_border = Border(
            left=Side(style='thin'), right=Side(style='thin'),
            top=Side(style='thin'), bottom=Side(style='thin')
        )
        wrap_align = Alignment(wrap_text=True, vertical="top")
        
        ws = wb.active
        ws.title = "批量比对汇总"
        
        summary_headers = ['序号', '对比合同', '整体相似度', '条款总数', '新增条款', '删除条款', '修改条款', '未变化条款']
        for col, h in enumerate(summary_headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border
        
        for row_idx, result in enumerate(batch_results, 2):
            row_data = [
                row_idx - 1,
                result.get('contract_name', '未知'),
                f"{result.get('overall_similarity', 0):.2%}",
                result.get('total_clauses_b', 0),
                result.get('stats', {}).get('added', 0),
                result.get('stats', {}).get('deleted', 0),
                result.get('stats', {}).get('modified', 0),
                result.get('stats', {}).get('unchanged', 0),
            ]
            for col, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.border = thin_border
                cell.alignment = wrap_align
        
        ws.column_dimensions['A'].width = 6
        ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 12
        ws.column_dimensions['D'].width = 10
        ws.column_dimensions['E'].width = 10
        ws.column_dimensions['F'].width = 10
        ws.column_dimensions['G'].width = 10
        ws.column_dimensions['H'].width = 10
        
        for result in batch_results:
            contract_name = result.get('contract_name', '未知')[:31]
            detail_data = []
            for idx, item in enumerate(result.get('diff_results', [])):
                if item.get('type') != 'unchanged':
                    detail_data.append({
                        '序号': idx + 1,
                        '类别': item.get('category', ''),
                        '类型': self._get_type_label(item.get('type', '')),
                        '相似度': f"{item.get('similarity', 0):.2%}",
                        '基准合同': item.get('old_text', ''),
                        '对比合同': item.get('new_text', ''),
                    })
            
            if detail_data:
                ws_detail = wb.create_sheet(contract_name)
                detail_headers = ['序号', '类别', '类型', '相似度', '基准合同', '对比合同']
                for col, h in enumerate(detail_headers, 1):
                    cell = ws_detail.cell(row=1, column=col, value=h)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = header_align
                    cell.border = thin_border
                
                type_fills = {
                    'added': PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid"),
                    'deleted': PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid"),
                    'modified': PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"),
                }
                
                for r_idx, d in enumerate(detail_data, 2):
                    fill = type_fills.get(list(result.get('diff_results', []))[d['序号']-1].get('type', '')) if d['序号']-1 < len(result.get('diff_results', [])) else None
                    for col, key in enumerate(detail_headers, 1):
                        cell = ws_detail.cell(row=r_idx, column=col, value=d.get(key, ''))
                        cell.border = thin_border
                        cell.alignment = wrap_align
                        if fill:
                            cell.fill = fill
                
                ws_detail.column_dimensions['A'].width = 6
                ws_detail.column_dimensions['B'].width = 12
                ws_detail.column_dimensions['C'].width = 8
                ws_detail.column_dimensions['D'].width = 10
                ws_detail.column_dimensions['E'].width = 40
                ws_detail.column_dimensions['F'].width = 40
        
        wb.save(filepath)
    
    def _export_batch_summary_pandas(self, batch_results, base_contract_name, filepath):
        summary_data = []
        for result in batch_results:
            summary_data.append({
                '对比合同': result['contract_name'],
                '整体相似度': f"{result['overall_similarity']:.2%}",
                '条款总数': result['total_clauses_b'],
                '新增条款': result['stats']['added'],
                '删除条款': result['stats']['deleted'],
                '修改条款': result['stats']['modified'],
                '未变化条款': result['stats']['unchanged']
            })
        
        df_summary = pd.DataFrame(summary_data)
        
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            df_summary.to_excel(writer, sheet_name='批量比对汇总', index=False)
            
            for result in batch_results:
                detail_data = []
                for idx, item in enumerate(result['diff_results']):
                    if item['type'] != 'unchanged':
                        detail_data.append({
                            '序号': idx + 1,
                            '类别': item['category'],
                            '类型': self._get_type_label(item['type']),
                            '相似度': f"{item['similarity']:.2%}",
                            '基准合同': item['old_text'],
                            '对比合同': item['new_text']
                        })
                
                if detail_data:
                    df_detail = pd.DataFrame(detail_data)
                    sheet_name = result['contract_name'][:31]
                    df_detail.to_excel(writer, sheet_name=sheet_name, index=False)
    
    def _export_batch_summary_csv(self, batch_results, base_contract_name, filepath):
        import csv
        csv_path = filepath.replace('.xlsx', '.csv')
        with open(csv_path, 'w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            writer.writerow(['对比合同', '整体相似度', '新增条款', '删除条款', '修改条款', '未变化条款'])
            for result in batch_results:
                writer.writerow([
                    result.get('contract_name', ''),
                    f"{result.get('overall_similarity', 0):.2%}",
                    result.get('stats', {}).get('added', 0),
                    result.get('stats', {}).get('deleted', 0),
                    result.get('stats', {}).get('modified', 0),
                    result.get('stats', {}).get('unchanged', 0),
                ])
        return csv_path, os.path.basename(csv_path)
