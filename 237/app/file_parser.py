import os
import re
from docx import Document
from app.config import ALLOWED_EXTENSIONS

def allowed_file(filename):
    return '.' in filename and \
           filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

def parse_txt(file_path):
    encodings = ['utf-8', 'gbk', 'gb2312', 'utf-16']
    for encoding in encodings:
        try:
            with open(file_path, 'r', encoding=encoding) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
        return f.read()

def parse_docx(file_path):
    doc = Document(file_path)
    full_text = []
    for para in doc.paragraphs:
        full_text.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    full_text.append(cell.text)
    return '\n'.join(full_text)

def parse_file(file_path):
    if not os.path.exists(file_path):
        raise FileNotFoundError(f"File not found: {file_path}")
    
    ext = file_path.rsplit('.', 1)[1].lower()
    
    if ext == 'txt':
        return parse_txt(file_path)
    elif ext == 'docx':
        return parse_docx(file_path)
    else:
        raise ValueError(f"Unsupported file format: {ext}")

def save_file(file_storage, upload_folder):
    if file_storage and allowed_file(file_storage.filename):
        filename = os.path.join(upload_folder, file_storage.filename)
        file_storage.save(filename)
        return filename
    return None
